#!/usr/bin/env python3
"""Доп. разведка: raw-формулы, гиперссылки, карта всех командных блоков на листе Showcase."""
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.formula import ArrayFormula

PATH = "/home/z/my-project/download/showcase_template.xlsx"
wb_f = load_workbook(PATH, data_only=False)
wb_v = load_workbook(PATH, data_only=True)
ws = wb_f["Showcase"]
wsv = wb_v["Showcase"]

def as_text(x):
    if isinstance(x, ArrayFormula):
        return f"ARRAY: {x.text}"
    return str(x) if x is not None else ""

print("=== 1. Все листы в книге ===")
print(wb_f.sheetnames)
print()

print("=== 2. RAW-формулы и значения в A4:S12 ===")
for r in range(4, 13):
    for c in range(1, 20):
        coord = f"{get_column_letter(c)}{r}"
        f, v = ws[coord].value, wsv[coord].value
        if isinstance(f, ArrayFormula) or (isinstance(f, str) and f.startswith("=")):
            print(f"  {coord}: FORMULA={as_text(f)!r} -> VALUE={v!r}")
print()

print("=== 3. Гиперссылки в ячейках A1:S12 ===")
for r in range(1, 13):
    for c in range(1, 20):
        coord = f"{get_column_letter(c)}{r}"
        cell = ws[coord]
        if cell.hyperlink is not None:
            print(f"  {coord}: target={cell.hyperlink.target!r}")
print()

print("=== 4. Карта командных блоков: поиск 'Team'/'Total DPR'/'Rot'/'Sheet' по всему листу ===")
markers = {}
for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
    for cell in row:
        v = cell.value
        if v is None:
            continue
        s = str(v).strip()
        if s in ("Team", "Total DPR", "Rot", "Sheet", "Damage", "Dist", "Field", "Extra Information"):
            markers.setdefault(s, []).append(cell.coordinate)
for k in ("Team", "Damage", "Dist", "Field", "Extra Information", "Total DPR", "Rot", "Sheet"):
    coords = markers.get(k, [])
    print(f"  {k!r}: {len(coords)} шт -> {coords}")
print()

print("=== 5. Колонка A (маркеры строк) ===")
for r in range(1, ws.max_row + 1):
    v = ws[f"A{r}"].value
    if v is not None and str(v).strip() != "":
        print(f"  A{r}: {v!r}")
print()

print("=== 6. Значения B4/D4/F4/H4 (кэш array-формул, data_only) ===")
for col in ("B", "D", "F", "H"):
    for r in range(4, 8):
        v = wsv[f"{col}{r}"].value
        f = ws[f"{col}{r}"].value
        if v is not None or f is not None:
            print(f"  {col}{r}: raw={as_text(f)[:60]!r} cached={v!r}")
print()

print("=== 7. Первая строка / заголовки колонок выше (строки 1-3, все 115 колонок) ===")
for r in range(1, 4):
    for c in range(1, ws.max_column + 1):
        coord = f"{get_column_letter(c)}{r}"
        v = ws[coord].value
        if v is not None and str(v).strip() != "" and not (isinstance(v, str) and v.startswith("=")):
            print(f"  {coord}: {str(v)[:60]!r}")
