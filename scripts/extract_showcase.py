#!/usr/bin/env python3
"""Извлечение шаблона входных данных A4:S12 из листа Showcase (Genshin Impact showcase table)."""
import json
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.formula import ArrayFormula

def as_text(x):
    """Преобразует значение ячейки в строку (обработка ArrayFormula)."""
    if x is None:
        return ""
    if isinstance(x, ArrayFormula):
        return f"AF({x.text})"
    if isinstance(x, str) and x.startswith("="):
        return x
    return str(x)

PATH = "/home/z/my-project/download/showcase_template.xlsx"
SHEET = "Showcase"

# Книга с формулами и книга с кэшированными значениями
wb_f = load_workbook(PATH, data_only=False)
wb_v = load_workbook(PATH, data_only=True)
ws_f = wb_f[SHEET]
ws_v = wb_v[SHEET]

print(f"=== Лист: {SHEET} | размер: {ws_f.dimensions} ===\n")

# Объединённые ячейки в диапазоне A1:S12
print("=== Merged cells (в пределах A1:S12) ===")
for mc in sorted(ws_f.merged_cells.ranges, key=lambda r: (r.min_row, r.min_col)):
    if mc.min_row <= 12 and mc.min_col <= 19:
        print(f"  {mc}")
print()

# Контекст: строки 1-3 (заголовки)
print("=== Контекст: строки 1-3 (столбцы A-S) ===")
for row in ws_f.iter_rows(min_row=1, max_row=3, min_col=1, max_col=19):
    for c in row:
        v = c.value
        if v is not None and str(v).strip() != "":
            print(f"  {c.coordinate}: {repr(v)}")
print()

# Извлечение A4:S12
def cell_info(coord):
    f = ws_f[coord].value
    v = ws_v[coord].value
    return f, v

def render(f, v):
    """Формула -> [значение]{формула}; иначе просто текст."""
    ft = as_text(f)
    if isinstance(f, ArrayFormula) or ft.startswith("=") or ft.startswith("AF("):
        vs = "" if v is None else str(v)
        return f"[{vs}]{ft}"
    return ft

print("=== ДИАПАЗОН A4:S12 (шаблон входных данных) ===")
header = [""] + [get_column_letter(i) for i in range(1, 20)]
print(" | ".join(f"{h:>18}" for h in header))
for r in range(4, 13):
    row_vals = [f"R{r}"]
    for col in range(1, 20):
        coord = f"{get_column_letter(col)}{r}"
        f, v = cell_info(coord)
        row_vals.append(f"{render(f, v)[:40]:>40}")
    print(" | ".join(row_vals))
print()

# Компактный дамп в JSON для последующего анализа
data = {}
for r in range(4, 13):
    row_data = {}
    for col in range(1, 20):
        coord = f"{get_column_letter(col)}{r}"
        f, v = cell_info(coord)
        if f is not None or v is not None:
            row_data[coord] = {"formula": as_text(f) if (isinstance(f, ArrayFormula) or as_text(f).startswith("=")) else None,
                                "value": v}
    if row_data:
        data[f"row_{r}"] = row_data

with open("/home/z/my-project/scripts/a4_s12_dump.json", "w", encoding="utf-8") as fh:
    json.dump(data, fh, ensure_ascii=False, indent=2, default=str)
print("JSON-дамп сохранён: /home/z/my-project/scripts/a4_s12_dump.json")
