#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
showcase_parser.py — универсальный парсер листов "Showcase" из калькуляторов
Genshin Impact (формат Timmie's Bird).

Шаблон командного блока (9 строк, якорь — ячейка 'Team' в J-колонке блока):
    строка +0:  [B..I]=карточки-картинки | J:K='Team' | L='Damage' | M='Dist' | N='Field' | P:S='Extra Information'
    строки +1..+4: J=имя персонажа (С-констелляция), L=DPR, M=доля урона, N=время на поле
    строка +4:  B/D/F/H = оружие персонажей 1-4
    строка +5:  B/D/F/H = сеты артефактов; J='Total DPR', L=суммарный DPR, M:N=дата (TODAY)
    строка +6:  B/D/F/H = требования ER; J=длина ротации (сек), L=DPS, M:N=автор
    строки +7..+8: B='Rot', C:N=текст ротации; B='Sheet' (ссылка на расчётный лист)
    B(строка+0) = массивная формула ='Имя листа'!$AA$256:$AM$262 -> имя расчётного листа

Сетка листа: блоки идут с шагом 19 колонок и 10 строк; слотов 6x6=36.

Поддерживаются источники: Google-URL (экспорт в xlsx или csv), локальный .xlsx,
локальный .csv. В CSV-режиме значения отформатированы ("1,959,082", "51.09%",
"14.5s", "9/17/2026") — парсер их нормализует; имена расчётных листов (calc_sheet)
в CSV недоступны (формулы теряются при экспорте).
"""
import csv
import io
import json
import os
import re
import sys
import urllib.request
from datetime import datetime, date, timedelta

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.formula import ArrayFormula

DEFAULT_SHEET_CANDIDATES = ("Showcase", "Showcase 2", "Showcase 3")


def download_google_xlsx(url_or_id, gid=None):
    """Скачивает Google-таблицу в xlsx. Возвращает bytes."""
    m = re.search(r"/spreadsheets/d/([a-zA-Z0-9-_]+)", url_or_id)
    sid = m.group(1) if m else url_or_id
    if gid is None:
        gm = re.search(r"[#&?]gid=(\d+)", url_or_id)
        gid = gm.group(1) if gm else None
    export = f"https://docs.google.com/spreadsheets/d/{sid}/export?format=xlsx"
    if gid:
        export += f"&gid={gid}"
    req = urllib.request.Request(export, headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(req, timeout=120) as resp:
        data = resp.read()
    if not data.startswith(b"PK"):
        raise ValueError(f"Google вернул не xlsx (нет доступа?): {data[:120]!r}")
    return data


def download_google_csv(url_or_id, gid=None):
    """Скачивает лист Google-таблицы в CSV (текст)."""
    m = re.search(r"/spreadsheets/d/([a-zA-Z0-9-_]+)", url_or_id)
    sid = m.group(1) if m else url_or_id
    if gid is None:
        gm = re.search(r"[#&?]gid=(\d+)", url_or_id)
        gid = gm.group(1) if gm else None
    url = f"https://docs.google.com/spreadsheets/d/{sid}/export?format=csv"
    if gid:
        url += f"&gid={gid}"
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(req, timeout=120) as resp:
        text = resp.read().decode("utf-8-sig", errors="replace")
    if not text or text.lstrip()[:1] == "<":
        raise ValueError(f"Google вернул не CSV (нет доступа?): {text[:120]!r}")
    return text


class _Cell:
    """Лёгкая замена ячейки openpyxl (для CSV-режима)."""
    __slots__ = ("value", "coordinate", "row", "column")

    def __init__(self, v, coordinate="", row=0, column=0):
        self.value = v
        self.coordinate = coordinate
        self.row = row
        self.column = column


class CsvSheet:
    """Адаптер CSV-текста под интерфейс листа openpyxl (координаты 1-based)."""

    def __init__(self, text, title="CSV"):
        self._rows = list(csv.reader(io.StringIO(text)))
        self.title = title

    @property
    def max_row(self):
        return len(self._rows)

    @property
    def max_column(self):
        return max((len(r) for r in self._rows), default=0)

    def cell(self, row, column):
        r = self._rows[row - 1] if 0 < row <= len(self._rows) else []
        v = r[column - 1] if 0 < column <= len(r) else None
        return _Cell(v, coordinate=f"{get_column_letter(column)}{row}", row=row, column=column)

    def iter_rows(self, min_row=1, max_row=None, min_col=1, max_col=None):
        max_row = self.max_row if max_row is None else max_row
        max_col = self.max_column if max_col is None else max_col
        for r in range(min_row, max_row + 1):
            yield [self.cell(r, c) for c in range(min_col, max_col + 1)]


def parse_calc_sheet_name(anchor_value):
    """Извлекает имя расчётного листа из массивной формулы ='Name'!$AA$256:$AM$262."""
    if anchor_value is None:
        return None
    text = anchor_value.text if isinstance(anchor_value, ArrayFormula) else str(anchor_value)
    m = re.match(r"^='?([^'!]+)'?!\$?A", text)
    return m.group(1) if m else None


def _num(v):
    """Число или None. Понимает отформатированные строки CSV:
    '1,959,082' -> 1959082.0; '51.09%' -> 0.5109; '14.5s' -> 14.5."""
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return v
    if isinstance(v, str):
        s = v.strip()
        if not s:
            return None
        pct = "%" in s
        s = s.replace(",", "").replace("%", "").rstrip("s").strip()
        try:
            x = float(s)
        except ValueError:
            return None
        return x / 100.0 if pct else x
    return None


def _iso(v):
    """ISO-дата из datetime / serial date / строки ('9/17/2026', '17.09.2026')."""
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, (int, float)) and not isinstance(v, bool) and 20000 < v < 80000:
        return (datetime(1899, 12, 30) + timedelta(days=int(v))).date().isoformat()
    if isinstance(v, str):
        s = v.strip()
        if not s:
            return None
        for fmt in ("%m/%d/%Y", "%d/%m/%Y", "%Y-%m-%d", "%Y/%m/%d", "%d.%m.%Y"):
            try:
                return datetime.strptime(s, fmt).date().isoformat()
            except ValueError:
                pass
        return s
    return None


def find_team_anchors(ws):
    """Находит все якоря 'Team' (верхняя-левая ячейка сводки блока)."""
    anchors = []
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.strip() == "Team":
                anchors.append((cell.row, cell.column))
    return sorted(anchors)


def parse_block(ws, wsf, r, c):
    """Разбирает один командный блок. (r, c) — позиция ячейки 'Team'.
    ws — лист с кэшированными значениями; wsf — лист с формулами (для имени расчётного листа)."""
    bc = c - 8  # колонка B блока (левый край карточек)
    out = {
        "_anchor": f"{get_column_letter(c)}{r}",
        "_range": f"{get_column_letter(bc)}{r}:{get_column_letter(c + 9)}{r + 8}",
    }

    # 1. Имя расчётного листа (из массивной формулы карточки — только из листа с формулами)
    out["calc_sheet"] = parse_calc_sheet_name(wsf.cell(row=r, column=bc).value)

    # 2. Персонажи (строки r+1..r+4): имя в col c, урон c+2, доля c+3, поле c+4
    chars = []
    for i in range(4):
        name = ws.cell(row=r + 1 + i, column=c).value
        if name is None or str(name).strip() in ("", "Total DPR"):
            break
        chars.append({
            "name": str(name).strip(),
            "damage": _num(ws.cell(row=r + 1 + i, column=c + 2).value),
            "damage_share": _num(ws.cell(row=r + 1 + i, column=c + 3).value),
            "field_time_s": _num(ws.cell(row=r + 1 + i, column=c + 4).value),
            "weapon": ws.cell(row=r + 4, column=bc + 2 * i).value,
            "artifact_set": ws.cell(row=r + 5, column=bc + 2 * i).value,
            "er_requirement": ws.cell(row=r + 6, column=bc + 2 * i).value,
        })
    out["characters"] = chars

    # 3. Итоги
    out["total_dpr"] = _num(ws.cell(row=r + 5, column=c + 2).value)
    out["dps"] = _num(ws.cell(row=r + 6, column=c + 2).value)
    out["rotation_time_s"] = _num(ws.cell(row=r + 6, column=c).value)
    out["date"] = _iso(ws.cell(row=r + 5, column=c + 3).value)
    out["author"] = ws.cell(row=r + 6, column=c + 3).value

    # 4. Ротация (текст в c-7 колонке на строке r+7)
    out["rotation"] = ws.cell(row=r + 7, column=bc + 1).value

    # 5. Доп. информация (merged P..S на строках r+1..r+8, якорь (r+1, c+6))
    extra = None
    for cc in range(c + 6, c + 10):
        for rr in range(r + 1, r + 9):
            v = ws.cell(row=rr, column=cc).value
            if v is not None and str(v).strip():
                extra = str(v)
                break
        if extra:
            break
    out["extra_info"] = extra

    # 6. Пустой ли слот (в CSV calc_sheet недоступен — ориентируемся на персонажей)
    out["is_empty"] = (not chars) or out["calc_sheet"] == "Sheet Name Here"
    return out


def _load_sheets(source, sheet_name=None, gid=None):
    """Универсальный загрузчик: Google-URL (xlsx/csv), локальный .xlsx, локальный .csv.
    Возвращает (лист значений, лист формул); для CSV оба — один адаптер CsvSheet."""
    if isinstance(source, str) and re.match(r"^https?://", source):
        if "format=csv" in source or "out:csv" in source:
            text = download_google_csv(source, gid=gid)
            g = CsvSheet(text, title="CSV (Google)")
            return g, g
        data = download_google_xlsx(source, gid=gid)
        wb = load_workbook(io.BytesIO(data), data_only=True)
        wbf = load_workbook(io.BytesIO(data), data_only=False)
    elif isinstance(source, str) and source.lower().endswith(".csv"):
        with open(source, encoding="utf-8-sig") as fh:
            g = CsvSheet(fh.read(), title=os.path.splitext(os.path.basename(source))[0])
        return g, g
    else:
        wb = load_workbook(source, data_only=True)
        wbf = load_workbook(source, data_only=False)

    # Выбор листа (только для xlsx-источников)
    if sheet_name:
        return wb[sheet_name], wbf[sheet_name]
    for cand in DEFAULT_SHEET_CANDIDATES:
        if cand in wb.sheetnames:
            return wb[cand], wbf[cand]
    raise ValueError(f"Лист Showcase не найден. Доступны: {wb.sheetnames}")


def parse_showcase(source, sheet_name=None, gid=None):
    """Главная функция: source = Google-URL (xlsx или csv), путь к .xlsx или путь к .csv."""
    ws, wsf = _load_sheets(source, sheet_name, gid)
    anchors = find_team_anchors(ws)
    blocks = [parse_block(ws, wsf, r, c) for r, c in anchors]
    filled = [b for b in blocks if not b["is_empty"]]
    return {
        "sheet": ws.title,
        "total_slots": len(blocks),
        "filled_slots": len(filled),
        "teams": blocks,
    }


def extract_range(source, cell_range="A4:S12", sheet_name=None, gid=None):
    """Сырое извлечение диапазона (для отладки/шаблона). Поддерживает xlsx и csv."""
    ws, _wsf = _load_sheets(source, sheet_name, gid)
    m = re.match(r"([A-Z]+)(\d+):([A-Z]+)(\d+)", cell_range)
    from openpyxl.utils import column_index_from_string
    c1, r1, c2, r2 = column_index_from_string(m.group(1)), int(m.group(2)), column_index_from_string(m.group(3)), int(m.group(4))
    grid = {}
    for row in ws.iter_rows(min_row=r1, max_row=r2, min_col=c1, max_col=c2):
        for cell in row:
            if cell.value is not None and str(cell.value).strip() != "":
                grid[cell.coordinate] = cell.value if not isinstance(cell.value, datetime) else cell.value.isoformat()
    return grid


if __name__ == "__main__":
    import argparse
    ap = argparse.ArgumentParser(description="Парсер Showcase-таблиц Genshin Impact")
    ap.add_argument("source", help="URL Google-таблицы или путь к файлу (.xlsx / .csv)")
    ap.add_argument("--gid", default=None)
    ap.add_argument("--sheet", default=None)
    ap.add_argument("--out", default=None)
    ap.add_argument("--range", default=None, help="сырое извлечение диапазона, напр. A4:S12")
    a = ap.parse_args()

    if a.range:
        res = extract_range(a.source, a.range, a.sheet, a.gid)
    else:
        res = parse_showcase(a.source, a.sheet, a.gid)

    text = json.dumps(res, ensure_ascii=False, indent=2, default=str)
    if a.out:
        with open(a.out, "w", encoding="utf-8") as f:
            f.write(text)
        print(f"Сохранено: {a.out}")
    else:
        print(text)
