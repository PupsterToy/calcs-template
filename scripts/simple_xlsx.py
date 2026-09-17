#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
simple_xlsx.py — ПРОСТОЙ парсер Showcase-листа из .xlsx (только xlsx).

Отличие от scripts/showcase_parser.py (тот остаётся рабочим "сложным" парсером):
  - читаются ТОЛЬКО видимые значения ячеек (кэшированные результаты формул);
  - НЕ читаются формулы, имена расчётных листов (calc_sheet), гиперссылки;
  - одна загрузка книги (data_only=True), без второго прохода по формулам.

Вход:  путь к .xlsx ИЛИ Google-ссылка на таблицу/лист (лист скачается в xlsx).
       --sheet ИМЯ  — выбрать лист по имени (по умолчанию: Showcase/Showcase 2/3)
Опции: --out файл.json — сохранить результат; иначе печать в консоль.

Формат командного блока (якорь — ячейка 'Team' в (r, c); bc = c - 8):
  r+1..r+4   col c      — имена персонажей ("C0 Sandrone")
              col c+2    — DPR (урон за ротацию)
              col c+3    — доля урона (0..1, в сумме 1)
              col c+4    — время на поле, с
  r+4  cols bc+0/2/4/6  — оружие 4 персонажей
  r+5  cols bc+0/2/4/6  — сеты артефактов; col c+2 — Total DPR; col c+3 — дата
  r+6  cols bc+0/2/4/6  — требования ER;   col c — длит. ротации; c+2 — DPS; c+3 — автор
  r+7  col bc+1          — текст ротации
  r+1..r+8, cols c+6..c+9 — примечания (Extra Information)

Сетка листа: слоты 6x6, шаг 19 колонок / 10 строк. Слот пуст ⇢ нет имён персонажей.
"""
import argparse
import io
import json
import re
import sys
import urllib.request
from datetime import datetime, date, timedelta

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

SHEET_CANDIDATES = ("Showcase", "Showcase 2", "Showcase 3")


# ---------------------------- вспомогательные ----------------------------

def download(url):
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(req, timeout=120) as resp:
        return resp.read()


def google_xlsx_url(src):
    """Строит URL экспорта листа Google-таблицы в xlsx."""
    m = re.search(r"/spreadsheets/d/([a-zA-Z0-9-_]+)", src)
    sid = m.group(1) if m else src
    url = f"https://docs.google.com/spreadsheets/d/{sid}/export?format=xlsx"
    gid = re.search(r"[#&?]gid=(\d+)", src)
    return url + (f"&gid={gid.group(1)}" if gid else "")


def text(v):
    """Строка без краёв или None."""
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def to_num(v):
    """Число или None. Понимает числа и отформатированные строки:
    '1,959,082' -> 1959082.0; '51.09%' -> 0.5109; '14.5s' -> 14.5."""
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = text(v)
    if not s:
        return None
    pct = "%" in s
    s = s.replace(",", "").replace("%", "").rstrip("s").strip()
    try:
        x = float(s)
    except ValueError:
        return None
    return x / 100.0 if pct else x


def to_date(v):
    """ISO-дата из datetime / serial-числа Excel / строки ('9/17/2026', ...)."""
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, (int, float)) and not isinstance(v, bool) and 20000 < v < 80000:
        return (datetime(1899, 12, 30) + timedelta(days=int(v))).date().isoformat()
    s = text(v)
    if not s:
        return None
    for fmt in ("%m/%d/%Y", "%d/%m/%Y", "%Y-%m-%d", "%Y/%m/%d", "%d.%m.%Y"):
        try:
            return datetime.strptime(s, fmt).date().isoformat()
        except ValueError:
            pass
    return s


# ---------------------------- разбор листа ----------------------------

def find_team_anchors(ws):
    """Все ячейки со значением 'Team' — левые верхние углы сводок блоков."""
    return sorted(
        (cell.row, cell.column)
        for row in ws.iter_rows()
        for cell in row
        if isinstance(cell.value, str) and cell.value.strip() == "Team"
    )


def parse_block(ws, r, c, row_index, col_index):
    """Один командный блок: только видимые значения."""
    bc = c - 8  # колонка B блока

    chars = []
    for i in range(4):
        name = text(ws.cell(row=r + 1 + i, column=c).value)
        if not name or name == "Total DPR":
            break
        chars.append({
            "name": name,
            "damage": to_num(ws.cell(row=r + 1 + i, column=c + 2).value),
            "damage_share": to_num(ws.cell(row=r + 1 + i, column=c + 3).value),
            "field_time_s": to_num(ws.cell(row=r + 1 + i, column=c + 4).value),
            "weapon": text(ws.cell(row=r + 4, column=bc + 2 * i).value),
            "artifact_set": text(ws.cell(row=r + 5, column=bc + 2 * i).value),
            "er_requirement": text(ws.cell(row=r + 6, column=bc + 2 * i).value),
        })

    extra = None
    for cc in range(c + 6, c + 10):
        for rr in range(r + 1, r + 9):
            v = text(ws.cell(row=rr, column=cc).value)
            if v:
                extra = v
                break
        if extra:
            break

    return {
        "slot": f"R{row_index[r]}·C{col_index[c]}",
        "anchor": f"{get_column_letter(c)}{r}",
        "range": f"{get_column_letter(bc)}{r}:{get_column_letter(c + 9)}{r + 8}",
        # имени расчётного листа в простом режиме нет — собираем из персонажей
        "team_name": " · ".join(re.sub(r"^C\d+\s+", "", ch["name"]) for ch in chars) or None,
        "characters": chars,
        "total_dpr": to_num(ws.cell(row=r + 5, column=c + 2).value),
        "dps": to_num(ws.cell(row=r + 6, column=c + 2).value),
        "rotation_time_s": to_num(ws.cell(row=r + 6, column=c).value),
        "rotation": text(ws.cell(row=r + 7, column=bc + 1).value),
        "extra_info": extra,
        "author": text(ws.cell(row=r + 6, column=c + 3).value),
        "date": to_date(ws.cell(row=r + 5, column=c + 3).value),
        "is_empty": not chars,
    }


def parse_showcase(ws, source):
    """Все блоки листа -> простой JSON-словарь."""
    anchors = find_team_anchors(ws)
    # Метки слотов считаем ПО ПОЗИЦИИ ЯКОРЯ в сетке (не по абсолютным координатам):
    # так парсер не зависит от лишних/пропущенных строк при экспорте.
    rows = sorted({r for r, _ in anchors})
    cols = sorted({c for _, c in anchors})
    row_index = {r: i + 1 for i, r in enumerate(rows)}
    col_index = {c: i + 1 for i, c in enumerate(cols)}

    teams = [parse_block(ws, r, c, row_index, col_index) for r, c in anchors]
    filled = [t for t in teams if not t["is_empty"]]
    return {
        "format": "showcase-simple/1",
        "source": source,
        "source_type": "xlsx",
        "sheet": ws.title,
        "extracted_at": datetime.now().isoformat(timespec="seconds"),
        "grid": {"rows": len(rows), "cols": len(cols)},
        "slots_total": len(teams),
        "slots_filled": len(filled),
        "teams": teams,
    }


def open_sheet(source, sheet_name=None):
    """Лист из локального .xlsx или Google-ссылки. Только значения."""
    if re.match(r"^https?://", source):
        data = download(google_xlsx_url(source))
        if not data.startswith(b"PK"):
            raise SystemExit(f"Google вернул не xlsx (нет доступа?): {data[:120]!r}")
        wb = load_workbook(io.BytesIO(data), data_only=True)
    else:
        # data_only=True: значения, которые отображаются на листе.
        # (если файл сохранён без кэша формул, такие ячейки дадут None)
        wb = load_workbook(source, data_only=True)
    if sheet_name:
        if sheet_name not in wb.sheetnames:
            raise SystemExit(f"Лист '{sheet_name}' не найден. Доступны: {wb.sheetnames}")
        return wb[sheet_name]
    for name in SHEET_CANDIDATES:
        if name in wb.sheetnames:
            return wb[name]
    raise SystemExit(f"Лист Showcase не найден. Доступны: {wb.sheetnames}")


def main():
    ap = argparse.ArgumentParser(
        description="Простой парсер Showcase-листа (xlsx): только видимые значения, без формул")
    ap.add_argument("source", help="путь к .xlsx или Google-ссылка на таблицу/лист")
    ap.add_argument("--sheet", default=None, help="имя листа (по умолчанию Showcase*)")
    ap.add_argument("--out", default=None, help="сохранить JSON в файл")
    a = ap.parse_args()

    ws = open_sheet(a.source, a.sheet)
    res = parse_showcase(ws, a.source)

    filled = [t for t in res["teams"] if not t["is_empty"]]
    nchars = sum(len(t["characters"]) for t in filled)
    print(f"Лист: {res['sheet']} | слотов: {res['slots_total']} | заполнено: {res['slots_filled']} "
          f"| персонажей: {nchars}")
    # контроль качества: суммы по каждой команде
    warns = 0
    for t in filled:
        d_sum = sum(c["damage"] or 0 for c in t["characters"])
        if t["total_dpr"] and abs(d_sum - t["total_dpr"]) / t["total_dpr"] > 0.005:
            print(f"  ! {t['slot']}: сумма урона {d_sum:.0f} != Total DPR {t['total_dpr']:.0f}", file=sys.stderr)
            warns += 1
    print(f"Проверка сумм: {'OK' if not warns else f'{warns} расхождений'}")

    out = json.dumps(res, ensure_ascii=False, indent=2)
    if a.out:
        with open(a.out, "w", encoding="utf-8") as f:
            f.write(out)
        print(f"Сохранено: {a.out}")
    else:
        print(out)


if __name__ == "__main__":
    main()
