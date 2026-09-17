#!/usr/bin/env python3
"""Проверка полной книги: карточки персонажей (источник array-формул), даты M9-аналогов, все формулы Showcase."""
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.formula import ArrayFormula

FULL = "/home/z/my-project/download/full_workbook.xlsx"

print("=== 1. Источник карточки: 'Sand Yae Ode Alyo'!AA256:AM262 ===")
wb1 = load_workbook(FULL, read_only=True, data_only=True)
ws1 = wb1["Sand Yae Ode Alyo"]
for row in ws1.iter_rows(min_row=255, max_row=263, min_col=26, max_col=40):
    for c in row:
        if c.value is not None and str(c.value).strip() != "":
            v = str(c.value).replace("\n", "\\n")[:80]
            print(f"  {c.coordinate}: {v!r}")
wb1.close()
print()

print("=== 2. Showcase в полной книге: все формулы (raw) ===")
wb2 = load_workbook(FULL, read_only=True, data_only=False)
ws2 = wb2["Showcase"]
count = 0
for row in ws2.iter_rows(min_row=1, max_row=65, max_col=115):
    for c in row:
        v = c.value
        if isinstance(v, ArrayFormula):
            count += 1
            if count <= 40:
                print(f"  {c.coordinate}: ARRAY {v.text[:90]}")
        elif isinstance(v, str) and v.startswith("="):
            count += 1
            if count <= 40:
                print(f"  {c.coordinate}: {v[:90]}")
print(f"  ... всего формул: {count}")
print()

print("=== 3. Аналоги M9 (даты) во всех блоках строки 4-12 ===")
wb3 = load_workbook(FULL, read_only=True, data_only=True)
ws3 = wb3["Showcase"]
# M-колонки блоков: M(13), AF(32), AY(51), BR(70), CK(89), DD(108)
cols = [13, 32, 51, 70, 89, 108]
for r in (9, 19, 29, 39, 49, 59):
    vals = []
    for c in cols:
        cell = ws3.cell(row=r, column=c).value
        vals.append(f"{get_column_letter(c)}{r}={cell}")
    print("  " + " | ".join(vals))
wb3.close()
